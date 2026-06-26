"""
端到端测试 generate_three_day_key_area.py 和 generate_ten_day_key_area.py：
  - 构造一份伪 240 XML（含东城/西城/天安门 3 个 station，每个 station 给 20 个时次数据）；
  - 强制走 local 数据源 + 离线 LLM 兜底；
  - 调用两个工具的核心方法生成文档（用 _build_table / _generate_summary 兜底分支即可，无需 docx 上传链路）；
  - 把渲染后的文档结构 dump 出来检查。
"""
from __future__ import annotations

import importlib.util
import os
import sys
from datetime import datetime
from xml.etree.ElementTree import Element, SubElement, ElementTree

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TMP_DIR = os.path.join(REPO, "_inspect_tmp")
os.makedirs(TMP_DIR, exist_ok=True)


# ---------- 构造伪 240 XML ----------

# 起报时间：2026 年 06 月 09 日 11 时 → 基础日是 9 日；
# 时次 hour 取 12, 24, 36, ..., 240（共 20 个时次）
BASE_TS = "202606091100"
HOURS = list(range(12, 241, 12))  # 12..240 step 12 → 20 个值

# 每个区域给一组逐时次假数据：天气、气温、风向、风力
# 故意做一点差异以验证两个工具的分支
SAMPLE_DATA = {
    "东城": [
        # (wp, t, wdir, wspeed)
        ("多云", "26", "南", "1、2级"),  # 9 日白天
        ("晴", "16", "南转北", "1、2级"),  # 9 日夜间
        ("晴间多云", "28", "南", "1、2级"),
        ("多云", "18", "南转北", "1、2级"),
        ("多云", "29", "北转南", "2、3级"),  # 11 日白天
        ("阴", "19", "北", "2、3级"),
        ("雷阵雨", "27", "北", "3、4级"),  # 12 日白天 含降水
        ("多云", "20", "北", "2、3级"),
        ("多云转晴", "30", "南", "2、3级"),
        ("晴", "21", "南", "1、2级"),
        ("晴", "31", "南", "1、2级"),
        ("晴", "20", "南", "1、2级"),
        ("多云", "29", "南", "1、2级"),
        ("阴", "21", "南", "1、2级"),
        ("阵雨", "26", "北", "2、3级"),  # 16 日降水
        ("多云", "20", "北", "2、3级"),
        ("晴", "27", "南", "2、3级"),
        ("晴", "20", "南", "2、3级"),
        ("多云", "28", "南", "2、3级"),
        ("多云", "19", "南", "1、2级"),
    ],
    "西城": [
        ("晴", "27", "南", "2级"),
        ("晴", "15", "南转北", "2级"),
        ("晴", "29", "南", "2、3级"),
        ("晴间多云", "17", "北", "2、3级"),
        ("多云", "28", "北", "3级"),
        ("晴", "18", "北", "2、3级"),
        ("多云", "27", "北", "2、3级"),
        ("晴", "19", "北", "1、2级"),
        ("晴", "30", "南", "2级"),
        ("晴", "20", "南", "1级"),
        ("晴", "31", "南", "1级"),
        ("晴", "21", "南", "1级"),
        ("晴", "30", "南", "1级"),
        ("晴", "20", "南", "1级"),
        ("多云", "28", "北", "2、3级"),
        ("晴", "18", "北", "2级"),
        ("晴", "29", "南", "2级"),
        ("晴", "20", "南", "2级"),
        ("晴", "30", "南", "2级"),
        ("晴", "21", "南", "1级"),
    ],
    "天安门": [
        ("多云", "26", "南", "2级"),
        ("晴", "16", "南转北", "1、2级"),
        ("晴间多云", "28", "南", "2、3级"),
        ("晴", "18", "北", "1、2级"),
        ("多云", "28", "北", "3、4级"),
        ("多云转晴", "19", "北", "2、3级"),
        ("晴间多云", "27", "北", "2、3级"),
        ("晴", "20", "北", "1、2级"),
        ("多云", "29", "南", "2级"),
        ("晴", "20", "南", "1级"),
        ("晴", "30", "南", "1、2级"),
        ("晴", "21", "南", "1级"),
        ("多云", "29", "南", "1级"),
        ("晴", "20", "南", "1级"),
        ("多云", "28", "北", "2、3级"),
        ("晴", "19", "北", "1、2级"),
        ("晴", "29", "南", "2级"),
        ("晴", "20", "南", "1级"),
        ("多云", "30", "南", "2级"),
        ("晴", "20", "南", "1级"),
    ],
}


def build_fake_xml() -> str:
    out_dir = os.path.join(TMP_DIR, "fake_bj240")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"MSP2_BJ-MO_WF_ME_LNO_BJ_{BASE_TS}_00000-24012.xml")

    root = Element("root")
    stations = SubElement(root, "stations")
    for name, series in SAMPLE_DATA.items():
        station = SubElement(stations, "station", attrib={"stationname": name})
        for h, (wp, t, wdir, wspeed) in zip(HOURS, series):
            d = SubElement(station, "data")
            SubElement(d, "hour").text = str(h)
            SubElement(d, "wp").text = wp
            SubElement(d, "t").text = t
            SubElement(d, "wdir").text = wdir
            SubElement(d, "wspeed").text = wspeed

    ElementTree(root).write(out_path, encoding="utf-8", xml_declaration=True)
    return out_dir


# ---------- 动态加载工具模块 ----------

def load_tool(module_name: str, py_path: str):
    spec = importlib.util.spec_from_file_location(module_name, py_path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def run_one(tool_label: str, py_path: str, template_path: str, area: str, hour_limit: int, take_n: int) -> None:
    """直接调用 tool 内部已经实现的解析 + 表格 + 兜底概述 + docx 渲染，
    不走 FastAPI 上传链路（避免依赖 open_webui 内部组件）。"""
    print(f"\n========== {tool_label}（{area}） ==========")
    fake_dir = build_fake_xml()
    mod = load_tool(f"_t_{tool_label}", py_path)
    t = mod.Tools()
    t.valves.source_mode = "local"
    t.valves.xml_dir = fake_dir
    t.valves.template_path = template_path

    xml_path = t._find_latest_xml()
    assert xml_path, "没有找到伪 XML"
    base_time = t._parse_base_time(xml_path)
    area_key = t._normalize_area(area)
    print(f"area_key = {area_key}")

    data_list, available, matched = t._parse_xml(xml_path, area_key)
    assert data_list is not None, f"没匹配到 station，可用: {available}"
    print(f"matched_station = {matched!r}, available = {sorted(available)}")

    forecast = [d for d in data_list if d["hour"] <= hour_limit][:take_n]
    print(f"forecast 时次数 = {len(forecast)}")

    table = t._build_table(forecast, base_time)
    print(f"table 行数 = {len(table)}")
    for r in table[:3] + table[-2:]:
        print(" ", r)

    # 用本地兜底跑一遍概述（不调 LLM）
    stats = t._compute_stats(forecast, base_time)
    if tool_label.startswith("ten"):
        from datetime import timedelta as _td
        start_dt = base_time + _td(hours=12)
        end_dt = base_time + _td(hours=240)
        date_range_str = f"{start_dt.month}月{start_dt.day}日至{end_dt.month}月{end_dt.day}日"
        summary = t._fallback_summary(area_key, stats, date_range_str)
    else:
        summary = t._fallback_summary(area_key, stats)
    print(f"summary 兜底文案 = {summary}")

    # 渲染 docx
    from docxtpl import DocxTemplate
    from docx import Document
    doc = DocxTemplate(template_path)
    now = datetime.now()
    ctx = {
        "district_title": mod.AREA_DISPLAY_TITLE[area_key],
        "report_datetime": f"{now.year}年{now.month:02d}月{now.day:02d}日{now.hour:02d}时",
        "summary": summary,
        "table": table,
    }
    doc.render(ctx)
    t._merge_date_cells(doc.docx, table)

    out_path = os.path.join(TMP_DIR, f"e2e_{tool_label}_{area_key}.docx")
    doc.save(out_path)
    print(f"已保存: {out_path}")

    # 读取检查
    d2 = Document(out_path)
    print(f"--- 渲染后段落 ---")
    for i, p in enumerate(d2.paragraphs):
        print(f"P{i}: {p.text!r}")
    tab = d2.tables[0]
    print(f"--- 渲染后表格 rows={len(tab.rows)} ---")
    for ri, row in enumerate(tab.rows):
        # 显示第一列是否被 vMerge 合并
        from docx.oxml.ns import qn
        first_tc = row._tr.findall(qn("w:tc"))[0]
        tcPr = first_tc.find(qn("w:tcPr"))
        vm = tcPr.find(qn("w:vMerge")) if tcPr is not None else None
        vm_state = (vm.get(qn("w:val")) or "continue") if vm is not None else "-"
        print(f"R{ri} [vMerge={vm_state}]:", [c.text for c in row.cells])


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")

    THREE_PY = os.path.join(REPO, "tools", "generate_three_day_key_area.py")
    TEN_PY = os.path.join(REPO, "tools", "generate_ten_day_key_area.py")
    THREE_TPL = os.path.join(REPO, "weather_templates", "three_day_key_area", "template.docx")
    TEN_TPL = os.path.join(REPO, "weather_templates", "ten_day_key_area", "template.docx")

    for area in ["东城", "西城区", "天安门广场"]:
        run_one("three", THREE_PY, THREE_TPL, area, hour_limit=72, take_n=6)

    for area in ["东城", "西城区", "天安门"]:
        run_one("ten", TEN_PY, TEN_TPL, area, hour_limit=240, take_n=20)
